"""生成 T003/T012 二进制测试夹具（文本 PDF、扫描件、DOCX/XLSX）。

与 SDD 文档生成脚本无关：本脚本只写出
    tests/fixtures/attachments/notice.pdf   两页文本层 PDF（第二页含表格）
    tests/fixtures/ocr/scanned_notice.png   图像式扫描件（带轻微旋转与噪声）
    tests/fixtures/ocr/scanned_notice.pdf   仅含图像的扫描件
    tests/fixtures/office/notice.docx       标题/段落/列表/表格
    tests/fixtures/office/notice.xlsx       两个内容 sheet 与一个空 sheet
全部内容为虚构夹具，使用系统 DejaVu 字体与固定噪声种子；重复运行字节一致。

用法（需要项目依赖）：uv run --locked python tools/make_fixture_binaries.py
"""

from __future__ import annotations

import hashlib
import random
import re
import struct
import zipfile
import zlib
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests" / "fixtures"

FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
)
LINES = ("OCR SAMPLE 123", "NOTICE-2026-09-10", "FIXTURE ONLY")
FONT_SIZE = 44
MARGIN = 30
LINE_GAP = 18


def render_text_image() -> tuple[list[int], int, int]:
    font_path = next((path for path in FONT_CANDIDATES if Path(path).is_file()), None)
    if font_path is None:
        raise SystemExit("未找到 DejaVuSans 字体，无法生成扫描件夹具")
    font = ImageFont.truetype(font_path, FONT_SIZE)
    probe = Image.new("L", (10, 10))
    draw = ImageDraw.Draw(probe)
    widths = [draw.textlength(line, font=font) for line in LINES]
    line_height = FONT_SIZE + LINE_GAP
    width = int(max(widths)) + MARGIN * 2
    height = line_height * len(LINES) + MARGIN
    canvas = Image.new("L", (width, height), 248)
    draw = ImageDraw.Draw(canvas)
    for index, line in enumerate(LINES):
        draw.text((MARGIN, MARGIN // 2 + index * line_height), line, font=font, fill=18)
    canvas = canvas.rotate(-0.4, resample=Image.BICUBIC, expand=False, fillcolor=248)
    noise = random.Random(20260911)
    pixels = canvas.tobytes()
    values = bytearray(pixels)
    for index, value in enumerate(values):
        if value > 200:
            roll = noise.random()
            if roll < 0.015:
                values[index] = 120 + noise.randrange(0, 80)
            elif roll < 0.04:
                values[index] = 200 + noise.randrange(0, 40)
    return list(values), width, height


def write_png(path: Path, pixels: list[int], width: int, height: int) -> None:
    raw = b"".join(
        b"\x00" + bytes(pixels[y * width:(y + 1) * width]) for y in range(height)
    )

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + tag
            + data
            + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0)
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


def write_image_pdf(path: Path, pixels: list[int], width: int, height: int) -> None:
    """A4 页面，仅嵌入灰度图像，无文本层。"""
    page_w, page_h, margin = 595, 842, 40
    draw_w = page_w - 2 * margin
    draw_h = draw_w * height / width
    if draw_h > page_h - 2 * margin:
        draw_h = page_h - 2 * margin
        draw_w = draw_h * width / height
    x = (page_w - draw_w) / 2
    y = (page_h - draw_h) / 2
    image_data = zlib.compress(bytes(pixels), 9)
    content = (
        f"q {draw_w:.2f} 0 0 {draw_h:.2f} {x:.2f} {y:.2f} cm /Im0 Do Q".encode("ascii")
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_w} {page_h}] "
            "/Resources << /XObject << /Im0 4 0 R >> >> /Contents 5 0 R >>"
        ).encode("ascii"),
        (
            f"<< /Type /XObject /Subtype /Image /Width {width} /Height {height} "
            f"/ColorSpace /DeviceGray /BitsPerComponent 8 /Filter /FlateDecode "
            f"/Length {len(image_data)} >>\nstream\n"
        ).encode("ascii")
        + image_data
        + b"\nendstream",
        f"<< /Length {len(content)} >>\nstream\n".encode("ascii")
        + content
        + b"\nendstream",
    ]
    _write_pdf(path, objects)


PAGE_ONE_LINES = (
    "Sample Notice - Public Knowledge Collection",
    "Fixture only. Not a real institution document.",
    "Issued: 2026-09-10    Ref: FIXTURE-2026-001",
)

TABLE_ROWS = (
    ("Item", "Quantity", "Amount"),
    ("Widgets", "12", "340.00"),
    ("Gadgets", "3", "125.50"),
)
TABLE_COLUMN_X = (60, 220, 330)
TABLE_START_Y = 720
TABLE_ROW_GAP = 20


def write_text_pdf(path: Path) -> None:
    """两页文本层 PDF：第一页段落，第二页标题与三列表格。"""
    page_one = ["BT /F1 14 Tf 60 760 Td 18 TL"]
    for line in PAGE_ONE_LINES:
        page_one.append(f"({line}) Tj T*")
    page_one.append("ET")
    page_one_content = "\n".join(page_one).encode("ascii")

    page_two = ["BT /F1 13 Tf 60 770 Td (Appendix A - Summary Table) Tj ET", "BT /F1 11 Tf"]
    for row_index, row in enumerate(TABLE_ROWS):
        y = TABLE_START_Y - row_index * TABLE_ROW_GAP
        for column_x, value in zip(TABLE_COLUMN_X, row):
            page_two.append(f"1 0 0 1 {column_x} {y} Tm ({value}) Tj")
    page_two.append("ET")
    page_two_content = "\n".join(page_two).encode("ascii")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 6 0 R] /Count 2 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(page_one_content)} >>\nstream\n".encode("ascii")
        + page_one_content
        + b"\nendstream",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 7 0 R >>"
        ),
        f"<< /Length {len(page_two_content)} >>\nstream\n".encode("ascii")
        + page_two_content
        + b"\nendstream",
    ]
    _write_pdf(path, objects)


def _write_pdf(path: Path, objects: list[bytes]) -> None:
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    path.write_bytes(bytes(out))

FIXED_TIMESTAMP = datetime(2026, 9, 10, 0, 0, 0, tzinfo=timezone.utc)

OFFICE_TABLE_ROWS = (
    ("Item", "Quantity", "Amount"),
    ("Widgets", "12", "340.00"),
    ("Gadgets", "3", "125.50"),
)


def write_docx(path: Path) -> None:
    """标题样式、段落、列表与表格的 DOCX 夹具。"""
    from docx import Document

    document = Document()
    document.core_properties.created = FIXED_TIMESTAMP
    document.core_properties.modified = FIXED_TIMESTAMP
    document.add_heading("Fixture Notice - Office Formats", level=1)
    document.add_paragraph("Fixture only. Not a real institution document.")
    document.add_heading("Scope", level=2)
    document.add_paragraph("Applies to fixture tests only.", style="List Bullet")
    document.add_paragraph("Second fixture step.", style="List Number")
    table = document.add_table(rows=3, cols=3)
    for row_index, row in enumerate(OFFICE_TABLE_ROWS):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.save(path)
    normalize_zip_timestamps(path)


def write_xlsx(path: Path) -> None:
    """含表头单位标注、数据行与空 sheet 的 XLSX 夹具。"""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.properties.created = FIXED_TIMESTAMP
    workbook.properties.modified = FIXED_TIMESTAMP
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Item", "Quantity（单位：件）", "Amount（单位：元）"])
    summary.append(["Widgets", 12, 340.0])
    summary.append(["Gadgets", 3, 125.5])
    sheet = workbook.create_sheet("Notes")
    sheet.append(["Note"])
    sheet.append(["Fixture only"])
    workbook.create_sheet("Empty")
    workbook.save(path)
    normalize_zip_timestamps(path)


def normalize_zip_timestamps(path: Path) -> None:
    """固定 OOXML 包内时间戳与 core.xml 的 modified，保证重复生成字节一致。"""
    fixed_modified = FIXED_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode("ascii")
    with zipfile.ZipFile(path) as source:
        entries = [
            (info.filename, source.read(info.filename), info.compress_type, info.external_attr)
            for info in source.infolist()
        ]
    with zipfile.ZipFile(path, "w") as target:
        for name, data, compress_type, external_attr in entries:
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    lambda match: match.group(1) + fixed_modified + match.group(2),
                    data,
                )
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = compress_type
            info.external_attr = external_attr
            target.writestr(info, data)


def main() -> None:
    pixels, width, height = render_text_image()
    targets = [
        (FIXTURES / "ocr" / "scanned_notice.png", None),
        (FIXTURES / "ocr" / "scanned_notice.pdf", None),
        (FIXTURES / "attachments" / "notice.pdf", "text"),
    ]
    for path, kind in targets:
        path.parent.mkdir(parents=True, exist_ok=True)
        if kind == "text":
            write_text_pdf(path)
        elif path.suffix == ".png":
            write_png(path, pixels, width, height)
        else:
            write_image_pdf(path, pixels, width, height)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        print(f"{path.relative_to(ROOT)}  {path.stat().st_size} bytes  sha256={digest}")
    (FIXTURES / "office").mkdir(parents=True, exist_ok=True)
    write_docx(FIXTURES / "office" / "notice.docx")
    write_xlsx(FIXTURES / "office" / "notice.xlsx")
    for path in (FIXTURES / "office" / "notice.docx", FIXTURES / "office" / "notice.xlsx"):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        print(f"{path.relative_to(ROOT)}  {path.stat().st_size} bytes  sha256={digest}")


if __name__ == "__main__":
    main()
