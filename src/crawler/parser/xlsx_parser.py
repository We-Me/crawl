"""XLS/XLSX 表格解析（T012）。

只读模式打开 XLSX：每个 sheet 对应一个表格块，sheet 顺序映射为 page_no（从 1
开始），保留表头、数据行与原始行号；表头文本中的（单位）等标注按原文提取到
units，不补造单位。空 sheet 记 empty 页状态，不静默丢弃。旧式二进制 XLS 见
legacy_parser.py。
"""

from __future__ import annotations

import io
import logging
import re
from typing import List, Optional

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from crawler.parser.parsed_page import PageStatus, ParsedBlock, ParsedPage

logger = logging.getLogger(__name__)

EXTRACTION_METHOD = "openpyxl_read_only"
UNIT_PATTERN = re.compile(r"[（(]\s*(?:单位|unit)\s*[:：]?\s*([^）)]+)[）)]")


class XlsxParseError(ValueError):
    """XLSX 无法打开或不是有效工作簿。"""


def parse_xlsx(content: bytes, page_url: str = "") -> ParsedPage:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise XlsxParseError(f"XLSX 无法解析：{exc}") from exc

    blocks: List[ParsedBlock] = []
    statuses: List[PageStatus] = []
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            rows_with_numbers = _sheet_rows(sheet)
            block = _sheet_block(sheet.title, rows_with_numbers, sheet_index)
            if block is None:
                statuses.append(
                    PageStatus(
                        page_no=sheet_index,
                        status="empty",
                        extraction_method=EXTRACTION_METHOD,
                    )
                )
                continue
            blocks.append(block)
            statuses.append(
                PageStatus(
                    page_no=sheet_index,
                    status="ok",
                    line_count=len(block.structured_data["rows"]),
                    extraction_method=EXTRACTION_METHOD,
                )
            )
    finally:
        workbook.close()

    metadata_missing = ["publication_date"]
    metadata_missing.extend(
        f"sheet_{status.page_no}_{status.status}"
        for status in statuses
        if status.status != "ok"
    )
    title = _title(blocks)
    if not title:
        metadata_missing.append("title")
    logger.debug("XLSX 解析完成 url=%s sheets=%d blocks=%d", page_url, len(statuses), len(blocks))
    return ParsedPage(
        title=title,
        full_text="\n".join(block.text for block in blocks if block.text),
        blocks=tuple(blocks),
        extraction_method=EXTRACTION_METHOD,
        metadata_missing=tuple(metadata_missing),
        page_status=tuple(statuses),
        page_count=len(statuses),
        canonical_url=page_url or None,
    )


def _sheet_rows(sheet) -> List:
    rows = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = ["" if value is None else str(value).strip() for value in row]
        if any(values):
            rows.append((row_number, values))
    return rows


def _sheet_block(sheet_name: str, rows_with_numbers: List, page_no: int) -> Optional[ParsedBlock]:
    if not rows_with_numbers:
        return None
    headers = rows_with_numbers[0][1]
    data_rows = [values for _, values in rows_with_numbers[1:]]
    row_numbers = [number for number, _ in rows_with_numbers[1:]]
    structured_data = {
        "sheet": sheet_name,
        "header_row": rows_with_numbers[0][0],
        "headers": headers,
        "rows": data_rows,
        "row_numbers": row_numbers,
        "caption": None,
    }
    units = {header: match.group(1).strip() for header in headers
             for match in [UNIT_PATTERN.search(header)] if match}
    if units:
        structured_data["units"] = units
    text = "\n".join(
        "\t".join(values) for _, values in rows_with_numbers
    )
    return ParsedBlock(block_type="table", text=text, page_no=page_no, structured_data=structured_data)


def _title(blocks: List[ParsedBlock]) -> str:
    for block in blocks:
        sheet_name = (block.structured_data or {}).get("sheet")
        if sheet_name:
            return str(sheet_name)
    return ""
